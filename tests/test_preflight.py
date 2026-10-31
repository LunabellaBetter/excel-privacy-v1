"""Independently authored V1 release privacy boundary regressions (2026-10-19)."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import sys

import openpyxl
from openpyxl.packaging.custom import CustomPropertyList, StringProperty
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from privacy_handler import excel_privacy_tool as engine
from privacy_handler.gui_controller import GUIWorkflowController
from privacy_handler.conflict_engine import compare_workbook_against_ledger
from privacy_handler.restore_engine import restore_workbook_v1


def source_file(folder):
    book = openpyxl.Workbook()
    book.active.append(['Record', 'Secret'])
    book.active.append(['fictional-001', 'SYNTHETIC_PRIVATE_SENTINEL'])
    book.properties.creator = 'SYNTHETIC_AUTHOR'
    book.custom_doc_props.append(StringProperty(name='PrivateMetadata', value='SYNTHETIC_PROPERTY'))
    destination = folder / 'fixture.xlsx'
    book.save(destination)
    return destination


def create_run(folder):
    return engine.mask_workbook_v1(source_file(folder), output_dir=folder / 'safe',
                                  private_key_dir=folder / 'local',
                                  text_columns=['Secret'], record_key=['Record'])


def test_final_zip_contains_neither_protected_value_nor_private_metadata(tmp_path):
    result = create_run(tmp_path)
    with ZipFile(result['masked_workbook_path']) as package:
        combined = b'\n'.join(package.read(p) for p in package.namelist())
    for sentinel in (b'SYNTHETIC_PRIVATE_SENTINEL', b'SYNTHETIC_PROPERTY', b'SYNTHETIC_AUTHOR'):
        assert sentinel not in combined
    assert result['secret_key_path'].read_bytes() not in combined


def test_metadata_exception_cannot_be_silently_ignored(monkeypatch):
    book = openpyxl.Workbook()
    book.custom_doc_props.append(StringProperty(name='XPT_RunId', value='synthetic-run'))
    def denied(*args, **kwargs):
        raise RuntimeError('SYNTHETIC_PRIVATE_ERROR')
    monkeypatch.setattr(CustomPropertyList, 'append', denied)
    with pytest.raises(engine.PrivacyToolError, match='METADATA_SANITIZATION_FAILED') as failure:
        engine.sanitize_workbook_metadata(book)
    assert 'SYNTHETIC_PRIVATE_ERROR' not in str(failure.value)


@pytest.mark.parametrize('phase', ['mask', 'restore'])
def test_metadata_failure_aborts_both_export_paths(tmp_path, monkeypatch, phase):
    result = create_run(tmp_path) if phase == 'restore' else None
    def denied(*args):
        raise engine.PrivacyToolError('METADATA_SANITIZATION_FAILED')
    monkeypatch.setattr(engine, 'sanitize_workbook_metadata', denied)
    output = tmp_path / 'failed'
    with pytest.raises(engine.PrivacyToolError, match='METADATA_SANITIZATION_FAILED'):
        if phase == 'mask':
            engine.mask_workbook_v1(source_file(tmp_path), output_dir=output,
                private_key_dir=tmp_path / 'local', text_columns=['Secret'], record_key=['Record'])
        else:
            restore_workbook_v1(result['masked_workbook_path'], result['ledger_enc_path'],
                result['secret_key_path'].read_bytes(), output_dir=output)
    assert not list(output.glob('*.xlsx'))


def test_serialized_audit_blocks_silent_metadata_cleanup_failure(tmp_path, monkeypatch):
    book = openpyxl.Workbook()
    book.properties.creator = 'SYNTHETIC_PRIVATE_AUTHOR'
    monkeypatch.setattr(engine, 'sanitize_workbook_metadata', lambda wb: None)
    output = tmp_path / 'should-not-exist.xlsx'
    with pytest.raises(engine.PrivacyToolError, match='METADATA_EXPORT_AUDIT_FAILED'):
        engine.save_workbook_checked(book, output)
    assert not output.exists()


@pytest.mark.parametrize('extension', ['.xlsm', '.XLSM', '.xls', '.csv'])
def test_unsupported_extensions_block_source_and_return(tmp_path, extension):
    original = source_file(tmp_path)
    fake = original.with_suffix(extension)
    fake.write_bytes(original.read_bytes())
    with pytest.raises(engine.PrivacyToolError):
        engine.mask_workbook_v1(fake, output_dir=tmp_path / 'safe', private_key_dir=tmp_path / 'local')
    controller = GUIWorkflowController(private_keys_dir=tmp_path / 'local')
    assert not controller.load_source_workbook(fake)[0]
    assert not controller.load_returned_workbook(fake)[0]
    assert not (tmp_path / 'safe').exists()


@pytest.mark.parametrize('part', ['xl/vbaProject.bin', 'xl/pivotCache/cache.xml',
    'xl/slicers/slicer.xml', 'xl/connections.xml', 'customXml/item.xml',
    'xl/embeddings/object.bin', 'xl/externalLinks/link.xml'])
def test_blocked_parts_rejected_even_under_xlsx_name(tmp_path, part):
    path = source_file(tmp_path)
    with ZipFile(path, 'a') as package:
        package.writestr(part, b'synthetic payload')
    with pytest.raises(engine.PrivacyToolError, match='高级对象'):
        engine.assert_supported_workbook(path)


def test_corrupt_package_is_not_treated_as_safe(tmp_path):
    path = tmp_path / 'corrupt.xlsx'
    path.write_bytes(b'not an Office package')
    with pytest.raises(engine.PrivacyToolError):
        engine.assert_supported_workbook(path)


def test_stale_compare_cannot_overwrite_ai_changed_protected_cell(tmp_path):
    result = create_run(tmp_path)
    key = result['secret_key_path'].read_bytes()
    previous = compare_workbook_against_ledger(result['masked_workbook_path'], result['ledger'])
    book = openpyxl.load_workbook(result['masked_workbook_path'])
    book.active['B2'] = 'AI_CHANGED_PROTECTED_VALUE'
    book.save(result['masked_workbook_path'])
    restored = restore_workbook_v1(result['masked_workbook_path'], result['ledger_enc_path'], key,
                                  output_dir=tmp_path / 'restored', change_list=previous)
    output = openpyxl.load_workbook(restored.restored_workbook_path)
    assert output.active['B2'].value == 'AI_CHANGED_PROTECTED_VALUE'
    assert restored.restored_cells_count == 0
    assert restored.skipped_conflicts_count + restored.skipped_unconfirmed_count == 1


def test_key_directory_cannot_be_inside_selected_share_directory(tmp_path):
    with pytest.raises(engine.PrivacyToolError, match='Private key directory'):
        engine.mask_workbook_v1(source_file(tmp_path), output_dir=tmp_path / 'share',
            private_key_dir=tmp_path / 'share' / 'keys', text_columns=['Secret'], record_key=['Record'])


def test_recalculation_flags_created_when_missing():
    from privacy_handler.numeric_masking import ensure_workbook_recalculation_properties
    book = openpyxl.Workbook()
    book.calculation = None
    ensure_workbook_recalculation_properties(book)
    assert book.calculation.fullCalcOnLoad and book.calculation.forceFullCalc


@pytest.mark.parametrize('phase', ['mask', 'compare', 'restore'])
def test_background_error_survives_deferred_gui_dispatch(tmp_path, monkeypatch, phase):
    from types import SimpleNamespace
    from privacy_app import PrivacyApp
    import privacy_app
    queue, received = [], []
    failure = engine.PrivacyToolError('METADATA_SANITIZATION_FAILED')
    def failed_operation():
        raise failure
    controller = SimpleNamespace(last_compare_summary=None,
        validate_mask_preconditions=lambda: (True, ''),
        execute_mask=failed_operation, execute_compare=failed_operation, execute_restore=failed_operation)
    app = SimpleNamespace(controller=controller,
        safe_dir_var=SimpleNamespace(get=lambda: str(tmp_path)),
        restored_dir_var=SimpleNamespace(get=lambda: str(tmp_path)),
        _set_busy=lambda *args: None, after=lambda delay, callback: queue.append(callback))
    for name in ['mask','compare','restore']:
        setattr(app, '_on_'+name+'_error', received.append)
    class ImmediateWorker:
        def __init__(self, target, **kwargs):self.target=target
        def start(self):self.target()
    monkeypatch.setattr(privacy_app.threading, 'Thread', ImmediateWorker)
    getattr(PrivacyApp, '_start_'+phase+'_task')(app)
    assert len(queue)==1
    queue.pop()()  # Happens after the except scope has ended, as in real Tk.
    assert received==[failure]


def test_mask_failure_releases_busy_state_and_disables_old_prompt(monkeypatch):
    from types import SimpleNamespace
    from privacy_app import PrivacyApp
    import privacy_app
    events=[]
    button=SimpleNamespace(config=lambda **kwargs: events.append(kwargs))
    app=SimpleNamespace(_set_busy=lambda *args: events.append(args),
                        btn_copy_prompt=button,btn_view_full_prompt=button)
    monkeypatch.setattr(privacy_app.messagebox,'showerror',lambda *args: events.append(args))
    PrivacyApp._on_mask_error(app,engine.PrivacyToolError('METADATA_SANITIZATION_FAILED'))
    assert events[0][0] is False
    assert events.count({'state':'disabled'})==2
