"""Regression gates for record keys protected after earlier fields."""

from decimal import Decimal
from pathlib import Path
import sys

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from privacy_handler.conflict_engine import ChangeStatus, compare_workbook_against_ledger
from privacy_handler.excel_privacy_tool import mask_workbook_v1
from privacy_handler.recovery_ledger import load_encrypted_ledger
from privacy_handler.restore_engine import analyze_restore_coverage, restore_workbook_v1
from privacy_handler.table_scope import SheetProtectionScope


def _mask(tmp_path: Path, scopes: list[SheetProtectionScope], workbook: openpyxl.Workbook):
    source = tmp_path / "source.xlsx"
    workbook.save(source)
    result = mask_workbook_v1(
        input_path=source,
        output_dir=tmp_path / "safe",
        private_key_dir=tmp_path / "private",
        sheet_scopes=scopes,
        numeric_scale_factor=Decimal("1.5"),
    )
    key = result["secret_key_path"].read_bytes()
    ledger = load_encrypted_ledger(result["ledger_enc_path"], key)
    return source, result["masked_workbook_path"], result["ledger_enc_path"], key, ledger


def _single_key_run(tmp_path: Path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Snapshot"
    sheet.append(["EmployeeID", "Name", "PrivateNote"])
    sheet.append(["E001", "Fictional A", "note 1"])
    sheet.append(["E002", "Fictional B", "note 2"])
    sheet.append(["E003", "Fictional C", "note 3"])
    return _mask(tmp_path, [SheetProtectionScope(
        sheet_name="Snapshot", text_columns=["A", "B", "C"], record_key="B"
    )], book)


def test_later_protected_single_key_clean_roundtrip(tmp_path: Path):
    source, masked, ledger_path, key, ledger = _single_key_run(tmp_path)
    book = openpyxl.load_workbook(masked)
    sheet = book["Snapshot"]
    assert len(ledger.entries) == 9
    for entry in ledger.entries:
        column = entry.protected_field_id.rsplit(":", 1)[-1]
        row = next(r for r in range(2, 5) if sheet[f"{column}{r}"].value == entry.get_decoded_masked_value())
        assert entry.record_identity_masked == [{
            "field": "B", "type": "str", "value": sheet[f"B{row}"].value,
        }]
    book.close()

    changes = compare_workbook_against_ledger(masked, ledger_path, key)
    coverage = analyze_restore_coverage(ledger, changes)
    assert (coverage.restorable_cells_count, coverage.unconfirmed_cells_count) == (9, 0)
    restored = restore_workbook_v1(masked, ledger_path, key, output_dir=tmp_path / "restored")
    assert restored.restored_cells_count == 9
    original = openpyxl.load_workbook(source)
    output = openpyxl.load_workbook(restored.restored_workbook_path)
    assert [[c.value for c in row] for row in output.active] == [
        [c.value for c in row] for row in original.active
    ]


def test_changed_protected_value_stays_changed_and_other_cells_restore(tmp_path: Path):
    _, masked, ledger_path, key, _ = _single_key_run(tmp_path)
    book = openpyxl.load_workbook(masked)
    book.active["C3"] = "AI edit"
    returned = tmp_path / "changed_value.xlsx"
    book.save(returned)
    changes = compare_workbook_against_ledger(returned, ledger_path, key)
    assert sum(c.status == ChangeStatus.PROTECTED_VALUE_CONFLICT for c in changes.cell_changes) == 1
    result = restore_workbook_v1(returned, ledger_path, key, output_dir=tmp_path / "restored")
    output = openpyxl.load_workbook(result.restored_workbook_path)
    assert (result.restored_cells_count, result.skipped_conflicts_count) == (8, 1)
    assert output.active["C3"].value == "AI edit"
    assert output.active["A3"].value == "E002"


@pytest.mark.parametrize("replacement", ["AI changed key", None])
def test_changed_or_deleted_key_blocks_only_affected_record(tmp_path: Path, replacement):
    _, masked, ledger_path, key, ledger = _single_key_run(tmp_path)
    book = openpyxl.load_workbook(masked)
    book.active["B3"] = replacement
    returned = tmp_path / "changed_key.xlsx"
    book.save(returned)
    changes = compare_workbook_against_ledger(returned, ledger_path, key)
    coverage = analyze_restore_coverage(ledger, changes)
    assert (coverage.restorable_cells_count, coverage.unconfirmed_cells_count) == (6, 3)
    assert any(s.change_type == ChangeStatus.DELETED_ROW for s in changes.structural_changes)
    result = restore_workbook_v1(returned, ledger_path, key, output_dir=tmp_path / "restored")
    output = openpyxl.load_workbook(result.restored_workbook_path)
    assert (result.restored_cells_count, result.skipped_unconfirmed_count) == (6, 3)
    assert output.active["A3"].value != "E002"
    assert output.active["B3"].value == replacement
    assert output.active["A2"].value == "E001"


@pytest.mark.parametrize("keys,protected", [
    (["A", "B"], ["C", "B", "D"]),
    (["B", "C"], ["A", "B", "C", "D"]),
])
def test_composite_key_uses_final_values_and_blocks_edits(tmp_path: Path, keys, protected):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Composite"
    sheet.append(["Month", "EmployeeID", "Team", "Note"])
    sheet.append(["2026-01", "E001", "T1", "one"])
    sheet.append(["2026-02", "E001", "T2", "two"])
    sheet.append(["2026-01", "E002", "T2", "three"])
    _, masked, ledger_path, key, ledger = _mask(tmp_path, [SheetProtectionScope(
        sheet_name="Composite", text_columns=protected, record_key=keys,
    )], book)
    clean = compare_workbook_against_ledger(masked, ledger_path, key)
    assert analyze_restore_coverage(ledger, clean).restorable_cells_count == len(ledger.entries)
    restored = restore_workbook_v1(masked, ledger_path, key, output_dir=tmp_path / "restored")
    assert restored.restored_cells_count == len(ledger.entries)

    returned = tmp_path / "changed_composite.xlsx"
    changed = openpyxl.load_workbook(masked)
    changed.active[f"{keys[-1]}3"] = "AI changed key"
    changed.save(returned)
    changes = compare_workbook_against_ledger(returned, ledger_path, key)
    coverage = analyze_restore_coverage(ledger, changes)
    assert coverage.unconfirmed_cells_count == len(protected)
    result = restore_workbook_v1(returned, ledger_path, key, output_dir=tmp_path / "restored")
    assert result.skipped_unconfirmed_count == len(protected)
    output = openpyxl.load_workbook(result.restored_workbook_path)
    assert output.active[f"{keys[-1]}3"].value == "AI changed key"


def test_multi_sheet_headers_duplicate_amounts_groups_and_analysis_formula(tmp_path: Path):
    book = openpyxl.Workbook()
    history = book.active
    history.title = "History"
    history.append(["Title"])
    history.append(["Month", "EmployeeID", "Name", "Salary"])
    history.append(["2026-01", "E001", "Fictional A", 100])
    history.append(["2026-02", "E001", "Fictional A", 120])
    snapshot = book.create_sheet("Snapshot")
    snapshot.append(["EmployeeID", "Name", "Salary"])
    snapshot.append(["E001", "Fictional A", 120])
    payroll = book.create_sheet("Payroll")
    payroll.append(["Title"])
    payroll.append(["Subtitle"])
    payroll.append(["Month", "EmployeeID", "Name", "D", "E", "F", "Amount", "H", "Amount"])
    payroll.append(["2026-01", "E001", "Fictional A", None, None, None, 100, None, 10])
    payroll.append(["2026-02", "E001", "Fictional A", None, None, None, 120, None, 12])
    scopes = [
        SheetProtectionScope("History", header_row=2, text_columns=["B", "C"], numeric_columns=["D"], record_key=["A", "B"], numeric_column_groups={"D": "shared"}),
        SheetProtectionScope("Snapshot", text_columns=["A", "B"], numeric_columns=["C"], record_key="A", numeric_column_groups={"C": "shared"}),
        SheetProtectionScope("Payroll", header_row=3, text_columns=["B", "C"], numeric_columns=["G", "I"], record_key=["A", "B"], numeric_column_groups={"G": "local", "I": "local"}),
    ]
    _, masked, ledger_path, key, ledger = _mask(tmp_path, scopes, book)
    fields = {p.protected_field_id: p for p in ledger.protected_fields}
    assert fields["Payroll:G"].numeric_analysis_group == fields["Payroll:I"].numeric_analysis_group
    assert fields["Payroll:G"].protected_field_id != fields["Payroll:I"].protected_field_id
    assert fields["History:D"].numeric_analysis_group == fields["Snapshot:C"].numeric_analysis_group

    returned = tmp_path / "analysis.xlsx"
    edited = openpyxl.load_workbook(masked)
    analysis = edited.create_sheet("Analysis")
    analysis["A1"] = "='Snapshot'!A2"
    edited.save(returned)
    changes = compare_workbook_against_ledger(returned, ledger_path, key)
    assert analyze_restore_coverage(ledger, changes).restorable_cells_count == len(ledger.entries)
    result = restore_workbook_v1(returned, ledger_path, key, output_dir=tmp_path / "restored")
    output = openpyxl.load_workbook(result.restored_workbook_path)
    assert result.restored_cells_count == len(ledger.entries)
    assert output.sheetnames[-1] == "Analysis"
    assert output["Analysis"]["A1"].value == "='Snapshot'!A2"
    assert (output["Payroll"]["G4"].value, output["Payroll"]["I4"].value) == (100, 10)


def test_numeric_key_processed_after_text_field_uses_final_scaled_key(tmp_path: Path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "NumericKey"
    sheet.append(["PrivateNote", "RecordNumber", "Amount"])
    sheet.append(["one", 101, 10])
    sheet.append(["two", 102, 20])
    _, masked, ledger_path, key, ledger = _mask(tmp_path, [SheetProtectionScope(
        sheet_name="NumericKey", text_columns=["A"], numeric_columns=["B", "C"], record_key="B",
        numeric_column_groups={"B": "late_key", "C": "amount"},
    )], book)
    masked_book = openpyxl.load_workbook(masked)
    a_entries = [e for e in ledger.entries if e.protected_field_id == "NumericKey:A"]
    assert [e.record_identity_masked[0]["value"] for e in a_entries] == [
        masked_book.active["B2"].value, masked_book.active["B3"].value,
    ]
    changes = compare_workbook_against_ledger(masked, ledger_path, key)
    assert analyze_restore_coverage(ledger, changes).restorable_cells_count == len(ledger.entries)
