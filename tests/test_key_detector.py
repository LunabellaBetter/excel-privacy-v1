"""Tests for Checkpoint 3: Record Key Detection Foundation and Statistical Analysis."""

from __future__ import annotations

import io
from pathlib import Path
import sys
import time
import pytest
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import privacy_handler.excel_privacy_tool as tool
from privacy_handler.key_detector import (
    KeyCandidateStats,
    WorkbookKeyAnalysisResult,
    WorkbookKeyCandidate,
    analyze_column_values,
    analyze_sheet_key_candidates,
    analyze_workbook_key_candidates,
    is_empty_value,
    is_formula_value,
)


def test_perfect_unique_column_is_strong_candidate() -> None:
    """1. Test that 100% unique non-empty column without formula is a Strong Key Candidate."""
    values = ["00000001", "00000002", "00000003"]
    stats = analyze_column_values(values, sheet_name="员工表", column_name="员工工号")
    assert stats.total_count == 3
    assert stats.non_empty_count == 3
    assert stats.null_count == 0
    assert stats.unique_count == 3
    assert stats.duplicate_distinct_values == 0
    assert stats.duplicate_rows == 0
    assert stats.contains_formula is False
    assert stats.uniqueness_ratio == 1.0
    assert stats.is_strong_candidate is True


def test_leading_zero_strings_preserved() -> None:
    """2. Test that leading zero strings are not converted to integers."""
    values = ["00000001", "00000002", "00000003"]
    stats = analyze_column_values(values, sheet_name="员工表", column_name="员工工号")
    assert stats.unique_count == 3
    # Check that '00000001' and '1' would be distinct
    values_with_int = ["00000001", 1]
    stats2 = analyze_column_values(values_with_int)
    assert stats2.unique_count == 2
    assert stats2.is_strong_candidate is True


def test_none_is_empty() -> None:
    """3. Test that None is treated as empty."""
    assert is_empty_value(None) is True
    stats = analyze_column_values(["A", None, "B"])
    assert stats.null_count == 1
    assert stats.non_empty_count == 2
    assert stats.is_strong_candidate is False


def test_empty_string_is_empty() -> None:
    """4. Test that empty string is treated as empty."""
    assert is_empty_value("") is True
    stats = analyze_column_values(["A", "", "B"])
    assert stats.null_count == 1
    assert stats.non_empty_count == 2
    assert stats.is_strong_candidate is False


def test_whitespace_only_is_empty() -> None:
    """5. Test that whitespace-only string is treated as empty."""
    assert is_empty_value("   ") is True
    assert is_empty_value("\t\n") is True
    stats = analyze_column_values(["A", "   ", "B"])
    assert stats.null_count == 1
    assert stats.non_empty_count == 2
    assert stats.is_strong_candidate is False


def test_zero_is_not_empty() -> None:
    """6. Test that 0, 0.0, and '0' are NOT treated as empty."""
    assert is_empty_value(0) is False
    assert is_empty_value(0.0) is False
    assert is_empty_value("0") is False
    assert is_empty_value("00000000") is False

    stats = analyze_column_values([0, "0", "00000000"])
    assert stats.null_count == 0
    assert stats.non_empty_count == 3
    assert stats.unique_count == 3
    assert stats.is_strong_candidate is True


def test_false_is_not_empty() -> None:
    """7. Test that False is NOT treated as empty."""
    assert is_empty_value(False) is False
    stats = analyze_column_values([False, True])
    assert stats.null_count == 0
    assert stats.non_empty_count == 2
    assert stats.unique_count == 2
    assert stats.is_strong_candidate is True


def test_duplicate_column_not_strong() -> None:
    """8. Test that duplicate values cause is_strong_candidate to be False."""
    values = ["0001", "0002", "0001"]
    stats = analyze_column_values(values)
    assert stats.total_count == 3
    assert stats.null_count == 0
    assert stats.unique_count == 2
    assert stats.duplicate_distinct_values == 1
    assert stats.duplicate_rows == 2
    assert stats.is_strong_candidate is False


def test_partial_uniqueness_not_strong() -> None:
    """9. Test that near-unique column (e.g. 999 unique out of 1000) is NOT a strong candidate."""
    values = [f"ID_{i:04d}" for i in range(999)] + ["ID_0000"]  # 1 duplicate
    assert len(values) == 1000
    stats = analyze_column_values(values)
    assert stats.total_count == 1000
    assert stats.unique_count == 999
    assert stats.uniqueness_ratio == 0.999
    assert stats.duplicate_distinct_values == 1
    assert stats.duplicate_rows == 2
    assert stats.is_strong_candidate is False


def test_all_empty_column_not_candidate() -> None:
    """10. Test that an all-empty column has 0 unique values and is not a candidate."""
    values = [None, "", "   ", None]
    stats = analyze_column_values(values)
    assert stats.total_count == 4
    assert stats.non_empty_count == 0
    assert stats.null_count == 4
    assert stats.unique_count == 0
    assert stats.uniqueness_ratio == 0.0
    assert stats.is_strong_candidate is False


def test_formula_column_not_strong_candidate() -> None:
    """11. Test that columns containing Excel formulas are never marked as strong candidates."""
    assert is_formula_value("=A1&B1") is True
    assert is_formula_value("NORMAL") is False

    values = ["=A1&B1", "=A2&B2", "=A3&B3"]
    stats = analyze_column_values(values)
    assert stats.total_count == 3
    assert stats.null_count == 0
    assert stats.unique_count == 3
    assert stats.contains_formula is True
    assert stats.formula_count == 3
    assert stats.is_strong_candidate is False


def test_same_key_across_multiple_sheets_is_not_cross_sheet_duplicate() -> None:
    """12. Test that the same key value appearing across multiple sheets is evaluated independently."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "人员历史"
    ws1.append(["员工工号", "部门"])
    ws1.append(["00001234", "研发部"])
    ws1.append(["00005678", "市场部"])

    ws2 = wb.create_sheet(title="人员快照")
    ws2.append(["员工工号", "职级"])
    ws2.append(["00001234", "P6"])  # Same key as in 人员历史
    ws2.append(["00005678", "P7"])

    result = analyze_workbook_key_candidates(wb)
    assert result.total_sheets == 2
    emp_candidate = next(c for c in result.candidates if c.logical_field == "员工工号")
    assert emp_candidate.present_sheet_count == 2
    assert emp_candidate.strong_sheet_count == 2
    assert emp_candidate.weak_sheet_count == 0
    assert emp_candidate.missing_sheet_count == 0
    assert emp_candidate.is_overall_strong is True
    assert "员工工号" in result.strong_candidates


def test_candidate_can_be_missing_from_some_sheets() -> None:
    """13. Test that a candidate column missing from some sheets reports correct coverage."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "人员总表"
    ws1.append(["员工工号", "姓名"])
    ws1.append(["E001", "Alice"])
    ws1.append(["E002", "Bob"])

    ws2 = wb.create_sheet(title="部门汇总")
    ws2.append(["部门名称", "预算"])
    ws2.append(["研发", 100000])
    ws2.append(["销售", 80000])

    result = analyze_workbook_key_candidates(wb)
    assert result.total_sheets == 2
    emp_cand = next(c for c in result.candidates if c.logical_field == "员工工号")
    assert emp_cand.present_sheet_count == 1
    assert emp_cand.strong_sheet_count == 1
    assert emp_cand.missing_sheet_count == 1
    assert emp_cand.weak_sheet_count == 0
    assert emp_cand.is_overall_strong is True


def test_workbook_summary_reports_present_strong_weak_missing() -> None:
    """14. Test comprehensive reporting of present, strong, weak, and missing sheets."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["ID", "Name"])
    ws1.append(["1", "A"])
    ws1.append(["2", "B"])

    ws2 = wb.create_sheet(title="Sheet2")
    ws2.append(["ID", "Name"])
    ws2.append(["1", "A"])
    ws2.append(["1", "B"])  # Duplicate in Sheet2!

    ws3 = wb.create_sheet(title="Sheet3")
    ws3.append(["OtherCol", "Desc"])
    ws3.append(["X", "Y"])  # ID is missing

    result = analyze_workbook_key_candidates(wb)
    id_cand = next(c for c in result.candidates if c.logical_field == "ID")
    assert id_cand.present_sheet_count == 2
    assert id_cand.strong_sheet_count == 1  # Strong in Sheet1
    assert id_cand.weak_sheet_count == 1    # Weak in Sheet2 due to duplicate
    assert id_cand.missing_sheet_count == 1  # Missing from Sheet3
    assert id_cand.is_overall_strong is False
    assert "ID" not in result.strong_candidates


def test_no_strong_candidate_is_valid_result() -> None:
    """15. Test that when no columns qualify as strong candidates, analysis returns cleanly without errors."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ColA", "ColB"])
    ws.append(["dup", ""])
    ws.append(["dup", ""])

    result = analyze_workbook_key_candidates(wb)
    assert isinstance(result, WorkbookKeyAnalysisResult)
    assert result.strong_candidates == []
    for cand in result.candidates:
        assert cand.is_overall_strong is False


def test_numeric_and_string_values_are_not_silently_normalized() -> None:
    """16. Test typed identity: 1234 (int), '1234' (str), and '001234' (str) remain 3 distinct keys."""
    values = [1234, "1234", "001234"]
    stats = analyze_column_values(values)
    assert stats.total_count == 3
    assert stats.null_count == 0
    assert stats.unique_count == 3
    assert stats.duplicate_distinct_values == 0
    assert stats.duplicate_rows == 0
    assert stats.is_strong_candidate is True


def test_key_detector_module_is_integrity_locked() -> None:
    """17. Test that key_detector.py is registered in LOCKED_FILES and checksum verification passes."""
    locked_names = [p.name for p in tool.LOCKED_FILES]
    assert "key_detector.py" in locked_names
    checksums = tool.current_checksums()
    assert "privacy_handler/key_detector.py" in checksums
    assert len(checksums["privacy_handler/key_detector.py"]) == 64
    tool.verify_self_before_sensitive()


def test_key_detector_performance_50k_rows_no_o_n_squared() -> None:
    """18. Benchmark: verify 50,000 rows x 10 columns completes in linear time with no O(N^2) explosion."""
    rows = 50000
    cols = 10

    # Build synthetic data in memory
    col_data = {f"Col_{c}": [f"VAL_{c}_{r:06d}" for r in range(rows)] for c in range(cols)}
    # Introduce duplicate and null in Col_0 and Col_1
    col_data["Col_0"][100] = col_data["Col_0"][0]
    col_data["Col_1"][200] = None

    t0 = time.perf_counter()
    stats_map: dict[str, KeyCandidateStats] = {}
    for col_name, values in col_data.items():
        stats_map[col_name] = analyze_column_values(values, sheet_name="BenchSheet", column_name=col_name)
    elapsed = time.perf_counter() - t0

    assert len(stats_map) == cols
    assert stats_map["Col_0"].is_strong_candidate is False
    assert stats_map["Col_1"].is_strong_candidate is False
    assert stats_map["Col_2"].is_strong_candidate is True
    # In pure Python hash table, 50k x 10 = 500k cells processes well under 2.0 seconds
    assert elapsed < 5.0, f"Benchmark took unexpectedly long: {elapsed:.2f}s"
