"""Tests for Checkpoint 2: Run Identity, FP8 Fingerprint, TokenAllocator, and Workbook XPT properties."""

from __future__ import annotations

import io
from pathlib import Path
import re
import sys
import uuid
import pytest
import openpyxl
from openpyxl.packaging.custom import StringProperty

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import privacy_handler.excel_privacy_tool as tool
from privacy_handler.identity import (
    AmbiguousRunIdentityError,
    IdentityError,
    TokenAllocator,
    compute_run_fingerprint,
    generate_run_id,
    generate_token,
    get_ledger_paths,
    is_masked_workbook,
    parse_token,
    read_workbook_identity,
    resolve_ledger_by_fingerprint,
    stamp_masked_workbook_identity,
)
from privacy_handler.recovery_ledger import (
    LedgerRunInfo,
    RecoveryLedger,
    save_encrypted_ledger,
)
from cryptography.fernet import Fernet


def test_run_id_is_uuid4() -> None:
    """1. Test that generate_run_id produces a valid 36-char standard UUIDv4 string."""
    run_id = generate_run_id()
    assert isinstance(run_id, str)
    assert len(run_id) == 36
    parsed = uuid.UUID(run_id, version=4)
    assert str(parsed) == run_id
    assert parsed.version == 4


def test_fp8_is_deterministic_for_same_run_id() -> None:
    """2. Test that FP8 fingerprint is strictly deterministic for identical run_id."""
    run_id = "4f3b7c21-9e12-4f80-87a6-2b1d7e824190"
    fp1 = compute_run_fingerprint(run_id)
    fp2 = compute_run_fingerprint(run_id)
    assert fp1 == fp2
    assert len(fp1) == 8
    assert fp1.isupper()
    assert re.match(r"^[A-F0-9]{8}$", fp1)


def test_fp8_differs_for_different_run_ids() -> None:
    """3. Test that FP8 differs for distinct run_ids (without claiming mathematical impossibility of collisions)."""
    run_a = generate_run_id()
    run_b = generate_run_id()
    assert run_a != run_b
    assert compute_run_fingerprint(run_a) != compute_run_fingerprint(run_b)


def test_token_format() -> None:
    """4. Test that generated token complies strictly with MSK-[RUN_FP8]-[TOKEN8]."""
    run_id = generate_run_id()
    fp8 = compute_run_fingerprint(run_id)
    token = generate_token(fp8)
    assert re.match(r"^MSK-[A-F0-9]{8}-[A-Z0-9]{8}$", token)
    parsed = parse_token(token)
    assert parsed is not None
    assert parsed[0] == fp8
    assert len(parsed[1]) == 8


def test_token_generation_has_no_collision_in_10000_samples() -> None:
    """5. Test that generating 10,000 tokens within a run has zero collisions."""
    run_id = generate_run_id()
    fp8 = compute_run_fingerprint(run_id)
    used_tokens: set[str] = set()
    tokens = [generate_token(fp8, used_tokens) for _ in range(10000)]
    assert len(tokens) == 10000
    assert len(set(tokens)) == 10000


def test_same_field_same_value_same_run_gets_same_token() -> None:
    """6. Test field-scoped consistency: same field + same value within run gets identical token."""
    allocator = TokenAllocator()
    t1 = allocator.allocate("员工工号", "00001234")
    t2 = allocator.allocate("员工工号", "00001234")
    assert t1 == t2


def test_same_field_same_value_different_run_gets_different_token() -> None:
    """7. Test run isolation: same field + same value in different runs gets different tokens."""
    alloc1 = TokenAllocator()
    alloc2 = TokenAllocator()
    t1 = alloc1.allocate("员工工号", "00001234")
    t2 = alloc2.allocate("员工工号", "00001234")
    assert t1 != t2


def test_different_fields_same_value_get_different_tokens() -> None:
    """8. Test field isolation: different fields with identical value get distinct tokens."""
    allocator = TokenAllocator()
    t_emp = allocator.allocate("员工工号", "1001")
    t_cost_center = allocator.allocate("成本中心", "1001")
    assert t_emp != t_cost_center


def test_explicit_linked_field_can_reuse_employee_id_token() -> None:
    """9. Test explicit linked field reuse: Key and Key2 reuse employee id tokens."""
    allocator = TokenAllocator()
    allocator.link_field("Key", "员工工号")
    allocator.link_field("Key2", "员工工号")

    t_emp = allocator.allocate("员工工号", "00001234")
    t_key = allocator.allocate("Key", "00001234")
    t_key2 = allocator.allocate("Key2", "00001234")

    assert t_emp == t_key == t_key2


def test_workbook_identity_properties_roundtrip() -> None:
    """10. Test that XPT custom doc props survive save and reload."""
    wb = openpyxl.Workbook()
    run_id = generate_run_id()
    stamp_masked_workbook_identity(wb, run_id, "1.0")

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    wb_loaded = openpyxl.load_workbook(bio)
    ident = read_workbook_identity(wb_loaded)
    assert ident["run_id"] == run_id
    assert ident["is_masked"] == "true"
    assert ident["protocol_version"] == "1.0"
    assert is_masked_workbook(wb_loaded) is True


def test_metadata_hygiene_removes_external_custom_props_but_preserves_xpt_props() -> None:
    """11. Key test: sanitize_workbook_metadata strips external properties but preserves XPT allowlist."""
    wb = openpyxl.Workbook()
    run_id = generate_run_id()
    stamp_masked_workbook_identity(wb, run_id, "1.0")

    # Add external sensitive properties
    wb.custom_doc_props.append(StringProperty(name="SensitiveOrg", value="InternalDeptSecret"))
    wb.custom_doc_props.append(StringProperty(name="AuthorCompany", value="ConfidentialInc"))

    tool.sanitize_workbook_metadata(wb)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    wb_loaded = openpyxl.load_workbook(bio)
    prop_names = {prop.name for prop in wb_loaded.custom_doc_props}

    # External properties must be wiped
    assert "SensitiveOrg" not in prop_names
    assert "AuthorCompany" not in prop_names

    # XPT properties must be preserved
    assert "XPT_RunId" in prop_names
    assert "XPT_Masked" in prop_names
    assert "XPT_Protocol" in prop_names

    ident = read_workbook_identity(wb_loaded)
    assert ident["run_id"] == run_id


def test_run_id_maps_to_ledger_filename() -> None:
    """12. Test that full run_id directly maps to <run_id>.recovery_ledger.enc."""
    run_id = generate_run_id()
    base_dir = Path("/mock/private_keys")
    ledger_path, key_path = get_ledger_paths(base_dir, run_id)
    assert ledger_path == base_dir / f"{run_id}.recovery_ledger.enc"
    assert key_path == base_dir / f"{run_id}.mask_secret.key"


def test_fp8_collision_is_treated_as_ambiguous(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """13. Test that when multiple ledgers produce the same FP8, resolver raises AmbiguousRunIdentityError."""
    ledger1 = tmp_path / "11111111-0000-0000-0000-000000000000.recovery_ledger.enc"
    ledger2 = tmp_path / "22222222-0000-0000-0000-000000000000.recovery_ledger.enc"
    ledger1.touch()
    ledger2.touch()

    # Synthetic simulation: mock compute_run_fingerprint to return identical FP8 for both candidates
    COLLISION_FP8 = "COLLIS01"

    def mock_fp(candidate_id: str) -> str:
        if candidate_id in ("11111111-0000-0000-0000-000000000000", "22222222-0000-0000-0000-000000000000"):
            return COLLISION_FP8
        return "OTHERFP8"

    monkeypatch.setattr("privacy_handler.identity.compute_run_fingerprint", mock_fp)

    with pytest.raises(AmbiguousRunIdentityError) as exc_info:
        resolve_ledger_by_fingerprint(tmp_path, COLLISION_FP8)

    assert COLLISION_FP8 in str(exc_info.value)
    assert "Ambiguous run fingerprint" in str(exc_info.value)

    # When single match exists, returns the matched path
    single_ledger = tmp_path / "33333333-0000-0000-0000-000000000000.recovery_ledger.enc"
    single_ledger.touch()
    monkeypatch.setattr(
        "privacy_handler.identity.compute_run_fingerprint",
        lambda cid: "UNIQUE01" if cid == "33333333-0000-0000-0000-000000000000" else "OTHERFP8",
    )
    resolved = resolve_ledger_by_fingerprint(tmp_path, "UNIQUE01")
    assert resolved == single_ledger

    # When 0 matches exist, raises FileNotFoundError
    with pytest.raises(FileNotFoundError):
        resolve_ledger_by_fingerprint(tmp_path, "NOTFOUND")


def test_identity_module_is_integrity_locked() -> None:
    """14. Test that identity.py is registered in LOCKED_FILES and checksum verification passes."""
    locked_names = [p.name for p in tool.LOCKED_FILES]
    assert "identity.py" in locked_names
    checksums = tool.current_checksums()
    assert "privacy_handler/identity.py" in checksums
    assert len(checksums["privacy_handler/identity.py"]) == 64
    tool.verify_self_before_sensitive()


