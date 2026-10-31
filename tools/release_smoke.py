"""Synthetic, platform-native frozen-app acceptance. Never uses business files."""
from pathlib import Path
import argparse
import hashlib
import json
import platform
import shutil
import subprocess
import sys
from zipfile import ZipFile

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles import Font


def make_fixture(path):
    book = openpyxl.Workbook()
    book.remove(book.active)
    layouts = [
        ('人员历史', 2, 60, ['年月','员工工号','姓名','部门','岗位','状态','地区','级别','序号','类别','标签','备注','标记','工龄年数','基本工资','奖金','加班费','加班小时','应出勤小时','实际出勤小时','补充','出生日期','手机号','邮箱','银行卡尾号','身份证后四位']),
        ('人员快照', 1, 20, ['员工工号','姓名','部门','岗位','地区','状态','级别','类别','年龄','当前月薪','奖金','标记','备注','假期余额','工时','补贴','序号','区域','渠道','标签','说明','邮箱','手机号']),
        ('薪酬明细', 3, 45, ['明细编号','员工工号','姓名','年月','类型','币种','金额','项目','金额','说明','备注','标记','银行卡尾号']),
    ]
    for name, header, count, columns in layouts:
        sheet = book.create_sheet(name)
        for _ in range(header-1):sheet.append(['Synthetic acceptance workbook'])
        sheet.append(columns)
        for i in range(count):
            emp = i % 20
            values = [f'fictional-{j}' for j in range(len(columns))]
            for j, label in enumerate(columns):
                if label == '员工工号':values[j] = f'V1SYNTH_PRIVATE_ID_{emp:03}'
                elif label == '姓名':values[j] = f'V1SYNTH_PRIVATE_NAME_{emp:03}'
                elif label == '年月':values[j] = f'2026-{i//20+1:02}'
                elif label == '明细编号':values[j] = f'DETAIL-{i:04}'
                elif label in ['金额','基本工资','奖金','加班费','当前月薪','工龄年数','加班小时','应出勤小时','实际出勤小时','年龄','假期余额','工时','补贴']:values[j] = (i+1)*(j+2)+0.25
                elif label in ['手机号','邮箱','银行卡尾号','身份证后四位','出生日期']:values[j] = f'V1SYNTH_PRIVATE_{j}_{emp:03}'
            sheet.append(values)
        sheet.freeze_panes = f'D{header+1}'
        sheet.auto_filter.ref = f'A{header}:{openpyxl.utils.get_column_letter(len(columns))}{header+count}'
        sheet.cell(header,1).font = Font(bold=True)
        sheet.cell(header,1).comment = Comment('Synthetic public instruction', 'Synthetic tester')
    hidden = book.create_sheet('隐藏配置')
    hidden.append(['Synthetic setting','no private content'])
    hidden.sheet_state = 'veryHidden'
    summary = book.create_sheet('原有汇总')
    summary.append(['序号','合计'])
    summary.append([1, "=SUM('薪酬明细'!G4:G48)"])
    chart = BarChart();chart.add_data(Reference(book['薪酬明细'],min_col=7,min_row=3,max_row=10),titles_from_data=True)
    summary.add_chart(chart,'D2')
    book.properties.creator = 'V1SYNTH_PRIVATE_AUTHOR'
    book.custom_doc_props.append(StringProperty(name='PrivateCustom',value='V1SYNTH_PRIVATE_CUSTOM'))
    path.parent.mkdir(parents=True,exist_ok=True)
    book.save(path)


def simulate_ai(masked, returned):
    book = openpyxl.load_workbook(masked)
    sheet = book['人员快照']
    rows = list(sheet.iter_rows(min_row=2, values_only=True))[::-1]
    for r, values in enumerate(rows,2):
        for c, value in enumerate(values,1):sheet.cell(r,c,value)
    analysis = book.create_sheet('AI分析')
    analysis.append(['员工引用','分析公式'])
    for r in range(2,22):analysis.append([f"='人员快照'!A{r}", f"='人员快照'!J{r}*2"])
    book.save(returned)


def launch(command, arguments):
    child = subprocess.run(command + arguments, capture_output=True, text=True, timeout=120)
    if child.returncode:
        raise RuntimeError(f'App subprocess failed: exit={child.returncode}; {child.stderr[:1000]}')


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--app',type=Path)
    parser.add_argument('--work-dir',type=Path,required=True)
    args=parser.parse_args()
    root=args.work_dir.resolve();root.mkdir(parents=True,exist_ok=True)
    command=[str(args.app.resolve())] if args.app else [sys.executable,str(Path(__file__).resolve().parents[1]/'privacy_app.py')]
    launch(command,['--verify-self'])
    launch(command,['--internal-ui-smoke',str(root/'tk.json')])
    source=root/'原始合成资料'/'复杂 HR 验收.xlsx';make_fixture(source)
    foreign=root/'无关文件.xlsx';make_fixture(foreign)
    launch(command,['--internal-smoke-mask',str(source),str(root/'本地恢复环境'),str(root/'mask.json')])
    masked=Path(json.loads((root/'mask.json').read_text())['masked_workbook_path'])
    with ZipFile(masked) as package:
        combined=b'\n'.join(package.read(p) for p in package.namelist())
    if b'V1SYNTH_PRIVATE_' in combined:raise RuntimeError('Protected sentinel remains in output package')
    returned=root/'AI 返回文件.xlsx';simulate_ai(masked,returned)
    # Separate process proves that recovery material survives an app restart.
    launch(command,['--internal-smoke-resume',str(returned),str(foreign),str(root/'本地恢复环境'),str(root/'resume.json')])
    resumed=json.loads((root/'resume.json').read_text())
    restored=openpyxl.load_workbook(resumed['restored_workbook_path'])
    original=openpyxl.load_workbook(source)
    for name, header in [('人员历史',2),('人员快照',1),('薪酬明细',3)]:
        before=list(original[name].iter_rows(min_row=header+1,values_only=True))
        after=list(restored[name].iter_rows(min_row=header+1,values_only=True))
        if sorted(map(repr,before))!=sorted(map(repr,after)):raise RuntimeError('Restored synthetic records differ from original')
    assert restored['AI分析']['B2'].value=="='人员快照'!J2*2"
    assert len(restored['原有汇总']._charts)==1
    assert restored['隐藏配置'].sheet_state=='veryHidden'
    assert restored.calculation.fullCalcOnLoad and restored.calculation.forceFullCalc
    report={'platform':platform.system(),'machine':platform.machine(),'python':platform.python_version(),
        'native_frozen_app':bool(args.app),'integrity':True,'fresh_process_recovery':True,
        'chinese_paths':True,'protected_sentinel_scan':True,'exact_protected_roundtrip':True,
        'ai_formula_preserved':True,'chart_preserved':True,'hidden_sheet_preserved':True,
        'recalculation_requested':True,'excel_2016_recalculation_executed':False,
        'total_protected_cells':resumed['total_protected_cells'],
        'restored_cells':resumed['restored_cells'], 'foreign_run_blocked':resumed['foreign_run_blocked']}
    report['native_tk'] = json.loads((root/'tk.json').read_text())
    (root/'acceptance.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))


if __name__=='__main__':main()
