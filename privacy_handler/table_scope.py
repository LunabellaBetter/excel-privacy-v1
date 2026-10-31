"""Worksheet scope, header row detection, column selection, and output paths foundation.

Provides unified business table coordinate system (header_row, data_start_row,
data_end_row), physical column identity preservation, duplicate header disambiguation,
empty row trimming, composite key validation, and safe output path resolution.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
import datetime
from decimal import Decimal
import os
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except Exception:  # pragma: no cover
    openpyxl = None
    get_column_letter = None

from privacy_handler.key_detector import is_empty_value, is_formula_value


# Header keyword set for recommending standard PII / sensitive data protection
RECOMMENDED_PROTECTION_KEYWORDS = frozenset({
    "姓名",
    "员工工号",
    "工号",
    "身份证",
    "身份证号",
    "证件号码",
    "出生日期",
    "手机号",
    "手机号码",
    "电话",
    "邮箱",
    "邮件地址",
    "家庭住址",
    "地址",
    "银行卡尾号",
    "银行卡号",
    "银行账号",
    "银行账户",
    "账户号",
})


class TableScopeError(Exception):
    """Base error for table scope and output path operations."""


@dataclass(frozen=True)
class ColumnDescriptor:
    """Descriptor for a physical column within a worksheet."""

    column_letter: str
    column_index: int  # 1-indexed
    raw_header: str | None
    header: str  # Display label or clean header name
    is_hidden: bool = False
    recommended_protection: bool = False
    is_numeric_candidate: bool = False

    @property
    def display_header(self) -> str:
        """Display label for UI or human report."""
        return self.header

    @property
    def is_blank_header(self) -> bool:
        """Whether the underlying cell had no meaningful header."""
        return self.raw_header is None or len(self.raw_header.strip()) == 0

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> ColumnDescriptor:
        raw_h = data.get("raw_header")
        return cls(
            column_letter=str(data["column_letter"]),
            column_index=int(data["column_index"]),
            raw_header=str(raw_h) if raw_h is not None else None,
            header=str(data["header"]),
            is_hidden=bool(data.get("is_hidden", False)),
            recommended_protection=bool(data.get("recommended_protection", False)),
            is_numeric_candidate=bool(data.get("is_numeric_candidate", False)),
        )


@dataclass(frozen=True)
class TableScope:
    """Coordinate system and boundary definition for a worksheet's business table."""

    sheet_name: str
    header_row: int
    data_start_row: int
    data_end_row: int | None
    selected_columns: list[ColumnDescriptor] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet_name": self.sheet_name,
            "header_row": self.header_row,
            "data_start_row": self.data_start_row,
            "data_end_row": self.data_end_row,
            "selected_columns": [col.to_dict() for col in self.selected_columns],
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> TableScope:
        return cls(
            sheet_name=str(data["sheet_name"]),
            header_row=int(data["header_row"]),
            data_start_row=int(data["data_start_row"]),
            data_end_row=int(data["data_end_row"]) if data.get("data_end_row") is not None else None,
            selected_columns=[
                ColumnDescriptor.from_dict(c) for c in data.get("selected_columns", [])
            ],
        )


@dataclass
class SheetProtectionScope:
    """Per-sheet configuration for multi-sheet masking runs."""

    sheet_name: str
    header_row: int = 1
    data_start_row: int | None = None
    text_columns: list[str] = field(default_factory=list)
    numeric_columns: list[str] = field(default_factory=list)
    record_key: list[str] | str | None = None
    numeric_analysis_group: str | None = None  # Legacy / default Group ID for sharing scale factor k
    numeric_column_groups: dict[str, str] = field(default_factory=dict)  # Mapping: column_letter -> effective group identity


@dataclass(frozen=True)
class OutputPaths:
    """Decoupled directory configuration for masked, restored, and private ledgers."""

    masked_output_dir: Path
    restored_output_dir: Path
    private_data_dir: Path

    def to_dict(self) -> dict[str, str]:
        return {
            "masked_output_dir": str(self.masked_output_dir),
            "restored_output_dir": str(self.restored_output_dir),
            "private_data_dir": str(self.private_data_dir),
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> OutputPaths:
        return cls(
            masked_output_dir=Path(data["masked_output_dir"]).resolve(),
            restored_output_dir=Path(data["restored_output_dir"]).resolve(),
            private_data_dir=Path(data["private_data_dir"]).resolve(),
        )


def _check_recommended_protection(header_text: str) -> bool:
    """Check if header matches recommended protection keywords."""
    clean = header_text.strip()
    if not clean or clean == "[空表头]":
        return False
    for kw in RECOMMENDED_PROTECTION_KEYWORDS:
        if kw in clean:
            return True
    return False


def extract_table_scope(
    ws: Any,
    header_row: int = 1,
    data_start_row: int | None = None,
    selected_column_letters: list[str] | None = None,
) -> TableScope:
    """Extract TableScope from an openpyxl worksheet.

    Inspects header row, column properties (hidden, candidate status),
    and determines non-empty data row boundary.
    """
    if header_row < 1:
        raise TableScopeError(f"header_row must be >= 1, got {header_row}")

    effective_data_start = header_row + 1 if data_start_row is None else data_start_row
    if effective_data_start <= header_row:
        raise TableScopeError(
            f"data_start_row ({effective_data_start}) must be strictly greater than header_row ({header_row})"
        )

    sheet_name = ws.title
    max_col = ws.max_column or 0
    max_row = ws.max_row or 0

    # If sheet has no non-empty cells anywhere (e.g. freshly created or blank sheet)
    # ws.max_column and ws.max_row in openpyxl default to 1 even when empty.
    sheet_is_empty = (
        max_col == 0
        or max_row < header_row
        or (hasattr(ws, "_cells") and len(ws._cells) == 0)
    )
    if sheet_is_empty:
        return TableScope(
            sheet_name=sheet_name,
            header_row=header_row,
            data_start_row=effective_data_start,
            data_end_row=None,
            selected_columns=[],
        )

    # Determine filter set of column letters if provided
    selected_filter = None
    if selected_column_letters is not None:
        selected_filter = {letter.upper().strip() for letter in selected_column_letters}

    # Helper for column hidden check
    def is_col_hidden(col_idx: int) -> bool:
        col_letter = get_column_letter(col_idx) if get_column_letter else str(col_idx)
        col_dim = ws.column_dimensions.get(col_letter)
        if col_dim is not None and col_dim.hidden:
            return True
        return False

    descriptors: list[ColumnDescriptor] = []

    # Sample rows for numeric candidate detection (sample up to 100 data rows)
    sample_rows = range(effective_data_start, min(max_row + 1, effective_data_start + 100))

    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx) if get_column_letter else str(col_idx)
        if selected_filter is not None and col_letter not in selected_filter:
            continue

        raw_val = ws.cell(row=header_row, column=col_idx).value
        raw_header_str: str | None = None
        if raw_val is not None:
            raw_header_str = str(raw_val)

        if is_empty_value(raw_val):
            clean_header = "[空表头]"
        else:
            clean_header = str(raw_val).strip()
            if not clean_header:
                clean_header = "[空表头]"

        hidden = is_col_hidden(col_idx)
        rec_prot = _check_recommended_protection(clean_header)

        # Detect if column contains numeric values in sample data rows (Strict Qualification):
        # Allow numeric protection ONLY if:
        # 1. Column has at least one non-empty value in sample rows.
        # 2. ALL non-empty static values in sample rows are genuine numbers (int, float, Decimal),
        #    and NOT bool, date, datetime, time, string (even if numeric like "001234"), or formula.
        has_numeric = False
        non_empty_count = 0
        all_numeric = True

        for r in sample_rows:
            v = ws.cell(row=r, column=col_idx).value
            if is_empty_value(v):
                continue
            non_empty_count += 1
            if is_formula_value(v):
                all_numeric = False
                break
            if isinstance(v, bool):
                all_numeric = False
                break
            if isinstance(v, (datetime.date, datetime.datetime, datetime.time)):
                all_numeric = False
                break
            if isinstance(v, (int, float, Decimal)):
                continue
            # If string or any other type
            all_numeric = False
            break

        if non_empty_count > 0 and all_numeric:
            has_numeric = True

        descriptors.append(
            ColumnDescriptor(
                column_letter=col_letter,
                column_index=col_idx,
                raw_header=raw_header_str,
                header=clean_header,
                is_hidden=hidden,
                recommended_protection=rec_prot,
                is_numeric_candidate=has_numeric,
            )
        )

    # Find actual data_end_row by scanning backwards from max_row to effective_data_start
    # Only rows with at least one non-empty cell in the selected columns count as data rows
    cols_to_check = [col.column_index for col in descriptors] if descriptors else list(range(1, max_col + 1))
    
    last_data_row: int | None = None
    for r in range(max_row, effective_data_start - 1, -1):
        row_has_data = False
        for c_idx in cols_to_check:
            val = ws.cell(row=r, column=c_idx).value
            if not is_empty_value(val):
                row_has_data = True
                break
        if row_has_data:
            last_data_row = r
            break

    return TableScope(
        sheet_name=sheet_name,
        header_row=header_row,
        data_start_row=effective_data_start,
        data_end_row=last_data_row,
        selected_columns=descriptors,
    )


def iter_table_scope_rows(ws: Any, scope: TableScope) -> list[dict[str, Any]]:
    """Iterate rows within the defined TableScope, returning typed row dictionaries.

    Each row dictionary contains keys formatted as '{col.column_letter}' or '{col.header}'.
    To ensure unambiguous referencing with duplicate headers, dictionary keys include:
    - col.column_letter (e.g. 'A', 'B')
    - col.header (if unique) or '{col.column_letter}｜{col.header}'
    - '__row__': physical 1-indexed row number in the worksheet.

    Rows beyond data_end_row (or trailing empty formatted rows) are excluded.
    """
    if scope.data_end_row is None or scope.data_end_row < scope.data_start_row:
        return []

    if not scope.selected_columns:
        return []

    # Check for duplicate headers
    header_counts: dict[str, int] = {}
    for col in scope.selected_columns:
        header_counts[col.header] = header_counts.get(col.header, 0) + 1

    rows: list[dict[str, Any]] = []
    for r in range(scope.data_start_row, scope.data_end_row + 1):
        row_dict: dict[str, Any] = {"__row__": r}
        has_any_data = False
        for col in scope.selected_columns:
            cell_val = ws.cell(row=r, column=col.column_index).value
            if not is_empty_value(cell_val):
                has_any_data = True

            # Populate by column_letter (physical identity)
            row_dict[col.column_letter] = cell_val

            # Populate by header identity
            if header_counts[col.header] > 1:
                # Disambiguate duplicate header with letter prefix
                row_dict[f"{col.column_letter}｜{col.header}"] = cell_val
                row_dict[f"{col.column_letter}_{col.header}"] = cell_val
            else:
                row_dict[col.header] = cell_val

        # Include row
        rows.append(row_dict)

    return rows


def validate_composite_key(
    rows: list[dict[str, Any]],
    key_fields: list[str],
) -> dict[str, Any]:
    """Validate user-selected composite key fields against loaded data rows.

    Requires:
    - total_count > 0
    - key_fields is non-empty
    - null_count == 0 (no row has empty/None/blank in any key component)
    - formula_count == 0 (no row has formula in any key component)
    - unique_tuple_count == total_count (strictly unique composite key tuples)
    - typed identity preserved ((type(v), v) per component tuple).

    Returns a validation result dictionary.
    """
    if not key_fields:
        return {
            "is_valid": False,
            "reason": "key_fields cannot be empty",
            "total_count": len(rows),
            "unique_count": 0,
            "null_count": 0,
            "formula_count": 0,
        }

    total_count = len(rows)
    if total_count == 0:
        return {
            "is_valid": False,
            "reason": "total_count is 0",
            "total_count": 0,
            "unique_count": 0,
            "null_count": 0,
            "formula_count": 0,
        }

    null_count = 0
    formula_count = 0
    seen_tuples: set[tuple[tuple[type, Any], ...]] = set()

    for row in rows:
        row_has_null = False
        row_has_formula = False
        tuple_components: list[tuple[type, Any]] = []

        for field_name in key_fields:
            if field_name not in row:
                raise TableScopeError(f"Field '{field_name}' not present in row keys")
            val = row[field_name]
            if is_empty_value(val):
                row_has_null = True
            if is_formula_value(val):
                row_has_formula = True

            # Typed identity preserved
            tuple_components.append((type(val), val))

        if row_has_null:
            null_count += 1
        if row_has_formula:
            formula_count += 1

        seen_tuples.add(tuple(tuple_components))

    unique_count = len(seen_tuples)
    is_valid = (
        null_count == 0
        and formula_count == 0
        and unique_count == total_count
    )

    reasons = []
    if null_count > 0:
        reasons.append(f"{null_count} rows have null/empty key components")
    if formula_count > 0:
        reasons.append(f"{formula_count} rows have formula values")
    if unique_count < total_count:
        reasons.append(f"composite key has duplicates ({unique_count} unique vs {total_count} total)")

    return {
        "is_valid": is_valid,
        "reason": "; ".join(reasons) if reasons else "valid",
        "total_count": total_count,
        "unique_count": unique_count,
        "null_count": null_count,
        "formula_count": formula_count,
        "key_fields": key_fields,
    }


def resolve_safe_output_path(
    base_dir: Path,
    source_path: Path,
    suffix: str = ".masked.xlsx",
) -> Path:
    """Resolve a collision-safe output file path within base_dir.

    Ensures:
    - Never overwrites the source file (source protection).
    - If a target file already exists, appends collision-safe suffix (_1, _2, etc.).
    - Validates directory writability pre-flight.
    - Preserves spaces, non-ASCII characters, and custom paths.
    """
    resolved_base = Path(base_dir).resolve()
    resolved_source = Path(source_path).resolve()

    if not resolved_base.exists():
        try:
            resolved_base.mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            raise TableScopeError(f"Cannot create output directory '{resolved_base}': {exc}") from exc

    if not os.access(resolved_base, os.W_OK):
        raise TableScopeError(f"Output directory '{resolved_base}' is not writable")

    stem = resolved_source.stem
    ext = resolved_source.suffix or ".xlsx"

    # Base candidate filename
    # E.g. stem + suffix (or if suffix contains extension, use stem + suffix)
    if suffix.startswith("."):
        candidate_name = f"{stem}{suffix}"
    else:
        candidate_name = f"{stem}.{suffix}"

    target_candidate = resolved_base / candidate_name

    # Collision avoidance loop
    counter = 1
    base_stem = target_candidate.stem
    target_ext = target_candidate.suffix

    while target_candidate == resolved_source or target_candidate.exists():
        candidate_name = f"{base_stem}_{counter}{target_ext}"
        target_candidate = resolved_base / candidate_name
        counter += 1

    return target_candidate
