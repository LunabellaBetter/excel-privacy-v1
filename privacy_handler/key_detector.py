"""Record Key (unique row linkage column) detection and statistical analysis.

Provides conservative column statistics, strict single-sheet and workbook-level
strong candidate evaluation, typed identity preservation, empty value normalization,
and formula exclusion for the V1 AI round-trip privacy protocol.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

try:
    import openpyxl
except Exception:  # pragma: no cover
    openpyxl = None


class KeyDetectorError(Exception):
    """Base error for key detector operations."""


def is_empty_value(val: Any) -> bool:
    """Determine if a cell value is considered empty.

    Empty values:
    - None
    - "" (empty string)
    - "   " (whitespace-only string)

    Non-empty values:
    - 0, 0.0 (numeric zero)
    - False (boolean false)
    - "0", "00000000" (string zero)
    """
    if val is None:
        return True
    if isinstance(val, str):
        return len(val.strip()) == 0
    return False


def is_formula_value(val: Any) -> bool:
    """Determine if a cell value represents an Excel formula."""
    if isinstance(val, str) and val.startswith("="):
        return True
    return False


@dataclass(frozen=True)
class KeyCandidateStats:
    """Statistical metrics and evaluation for a single column within a sheet."""

    sheet_name: str
    column_name: str
    total_count: int
    non_empty_count: int
    null_count: int
    unique_count: int
    duplicate_distinct_values: int
    duplicate_rows: int
    contains_formula: bool
    formula_count: int
    uniqueness_ratio: float
    is_strong_candidate: bool

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class WorkbookKeyCandidate:
    """Cross-sheet aggregation of candidate stats for a logical field."""

    logical_field: str
    present_sheet_count: int
    strong_sheet_count: int
    weak_sheet_count: int
    missing_sheet_count: int
    is_overall_strong: bool
    per_sheet_stats: dict[str, KeyCandidateStats]

    def to_dict(self) -> dict[str, Any]:
        return {
            "logical_field": self.logical_field,
            "present_sheet_count": self.present_sheet_count,
            "strong_sheet_count": self.strong_sheet_count,
            "weak_sheet_count": self.weak_sheet_count,
            "missing_sheet_count": self.missing_sheet_count,
            "is_overall_strong": self.is_overall_strong,
            "per_sheet_stats": {s: stats.to_dict() for s, stats in self.per_sheet_stats.items()},
        }


@dataclass(frozen=True)
class WorkbookKeyAnalysisResult:
    """Overall workbook-level key detection analysis result."""

    total_sheets: int
    sheet_names: list[str]
    candidates: list[WorkbookKeyCandidate]
    strong_candidates: list[str]

    def to_dict(self) -> dict[str, Any]:
        return {
            "total_sheets": self.total_sheets,
            "sheet_names": list(self.sheet_names),
            "candidates": [c.to_dict() for c in self.candidates],
            "strong_candidates": list(self.strong_candidates),
        }


def analyze_column_values(
    values: list[Any],
    sheet_name: str = "",
    column_name: str = "",
) -> KeyCandidateStats:
    """Analyze a list of column values, calculating conservative uniqueness and candidate status."""
    total_count = len(values)
    null_count = 0
    non_empty_count = 0
    formula_count = 0
    value_counts: dict[tuple[type, Any], int] = {}

    for v in values:
        if is_empty_value(v):
            null_count += 1
        else:
            non_empty_count += 1
            if is_formula_value(v):
                formula_count += 1
            try:
                key = (type(v), v)
            except TypeError:
                key = (type(v), str(v))
            value_counts[key] = value_counts.get(key, 0) + 1

    unique_count = len(value_counts)
    duplicate_distinct_values = sum(1 for c in value_counts.values() if c > 1)
    duplicate_rows = sum(c for c in value_counts.values() if c > 1)
    contains_formula = formula_count > 0
    uniqueness_ratio = (unique_count / non_empty_count) if non_empty_count > 0 else 0.0

    # Strong candidate rule: non-empty total, 0 nulls, 100% unique, no formulas
    is_strong_candidate = (
        total_count > 0
        and null_count == 0
        and unique_count == total_count
        and not contains_formula
    )

    return KeyCandidateStats(
        sheet_name=sheet_name,
        column_name=column_name,
        total_count=total_count,
        non_empty_count=non_empty_count,
        null_count=null_count,
        unique_count=unique_count,
        duplicate_distinct_values=duplicate_distinct_values,
        duplicate_rows=duplicate_rows,
        contains_formula=contains_formula,
        formula_count=formula_count,
        uniqueness_ratio=uniqueness_ratio,
        is_strong_candidate=is_strong_candidate,
    )


def safe_header_str(val: Any, col_idx: int) -> str:
    """Sanitize header value into non-empty string."""
    if val is None:
        return f"Column_{col_idx}"
    text = str(val).strip()
    if not text:
        return f"Column_{col_idx}"
    if "\r" in text or "\n" in text:
        text = text.replace("\r", " ").replace("\n", " ")
    return text


def analyze_sheet_key_candidates(
    ws: Any,
    header_row: int = 1,
    min_row: int | None = None,
    max_row: int | None = None,
) -> dict[str, KeyCandidateStats]:
    """Scan an openpyxl Worksheet and return stats for each column in the header row."""
    sheet_name = getattr(ws, "title", "Sheet")
    max_sheet_row = getattr(ws, "max_row", 0) or 0
    max_sheet_col = getattr(ws, "max_column", 0) or 0

    if max_sheet_row <= header_row or max_sheet_col == 0:
        # No data rows available
        return {}

    # Extract headers
    headers: list[tuple[str, int]] = []
    for col_idx in range(1, max_sheet_col + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        header_name = safe_header_str(cell.value, col_idx)
        headers.append((header_name, col_idx))

    start_row = min_row if min_row is not None else (header_row + 1)
    end_row = max_row if max_row is not None else max_sheet_row

    if end_row < start_row:
        return {}

    # Collect column values in a single linear pass over iter_rows
    col_data: dict[str, list[Any]] = {name: [] for name, _ in headers}
    col_idx_to_name: dict[int, str] = {idx: name for name, idx in headers}

    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=max_sheet_col):
        for cell in row:
            name = col_idx_to_name.get(cell.column)
            if name:
                # Retain raw value or formula representation
                raw_val = cell.value
                if cell.data_type == "f" and not (isinstance(raw_val, str) and raw_val.startswith("=")):
                    # Explicit formula type in openpyxl
                    raw_val = f"={raw_val}" if raw_val is not None else "="
                col_data[name].append(raw_val)

    # Analyze each column
    results: dict[str, KeyCandidateStats] = {}
    for name, values in col_data.items():
        results[name] = analyze_column_values(values, sheet_name=sheet_name, column_name=name)

    return results


def analyze_workbook_key_candidates(
    wb_or_path: Any,
    header_row: int = 1,
    target_sheets: list[str] | None = None,
) -> WorkbookKeyAnalysisResult:
    """Analyze all candidate record keys across sheets in a workbook."""
    close_needed = False
    if isinstance(wb_or_path, (str, Path)):
        if openpyxl is None:
            raise KeyDetectorError("openpyxl is required to load workbook files")
        wb = openpyxl.load_workbook(wb_or_path, data_only=False)
        close_needed = True
    else:
        wb = wb_or_path

    try:
        available_sheets = list(wb.sheetnames)
        if target_sheets is not None:
            sheets_to_process = [s for s in target_sheets if s in available_sheets]
        else:
            sheets_to_process = available_sheets

        total_sheets = len(sheets_to_process)
        sheet_results: dict[str, dict[str, KeyCandidateStats]] = {}
        all_fields_ordered: list[str] = []

        for sheet_name in sheets_to_process:
            ws = wb[sheet_name]
            stats_dict = analyze_sheet_key_candidates(ws, header_row=header_row)
            sheet_results[sheet_name] = stats_dict
            for field_name in stats_dict:
                if field_name not in all_fields_ordered:
                    all_fields_ordered.append(field_name)

        candidates: list[WorkbookKeyCandidate] = []
        for field_name in all_fields_ordered:
            present_sheets = [s for s in sheets_to_process if field_name in sheet_results[s]]
            present_count = len(present_sheets)
            missing_count = total_sheets - present_count

            per_sheet: dict[str, KeyCandidateStats] = {}
            strong_count = 0
            for s in present_sheets:
                col_stats = sheet_results[s][field_name]
                per_sheet[s] = col_stats
                if col_stats.is_strong_candidate:
                    strong_count += 1

            weak_count = present_count - strong_count
            is_overall_strong = (present_count > 0 and strong_count == present_count and weak_count == 0)

            candidates.append(
                WorkbookKeyCandidate(
                    logical_field=field_name,
                    present_sheet_count=present_count,
                    strong_sheet_count=strong_count,
                    weak_sheet_count=weak_count,
                    missing_sheet_count=missing_count,
                    is_overall_strong=is_overall_strong,
                    per_sheet_stats=per_sheet,
                )
            )

        strong_candidates = [c.logical_field for c in candidates if c.is_overall_strong]

        return WorkbookKeyAnalysisResult(
            total_sheets=total_sheets,
            sheet_names=sheets_to_process,
            candidates=candidates,
            strong_candidates=strong_candidates,
        )
    finally:
        if close_needed and hasattr(wb, "close"):
            wb.close()
