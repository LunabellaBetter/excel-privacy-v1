"""V1 desktop boundary and workbook I/O, independently implemented 2026-10-19.

Legacy upstream CLI commands are intentionally absent. The local V1 pipeline below
is retained from e70ee25; see docs/V1_RELEASE_CONTRACT.md for provenance and scope.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import zipfile
from typing import Any
from xml.etree import ElementTree as ET

from cryptography.fernet import Fernet
from openpyxl import load_workbook
from openpyxl.packaging.custom import CustomPropertyList

PROJECT_ROOT = Path(__file__).resolve().parent.parent
TOOL_DIR = PROJECT_ROOT / 'privacy_handler'
_MODULES = (
    'excel_privacy_tool', 'identity', 'key_detector', 'numeric_masking',
    'conflict_engine', 'restore_engine', 'prompt_generator', 'table_scope',
    'recovery_ledger', 'gui_controller', 'packaging_smoke',
)
LOCKED_FILES = [PROJECT_ROOT / 'privacy_app.py'] + [
    PROJECT_ROOT / 'privacy_handler' / (name + '.py') for name in _MODULES
] + [PROJECT_ROOT / 'privacy_handler' / 'README.md']
APPROVED_CHECKSUMS = PROJECT_ROOT / 'privacy_handler' / 'approved_checksums.json'
_METADATA_FIELDS = ('creator', 'lastModifiedBy', 'title', 'subject', 'description',
                    'keywords', 'category', 'contentStatus', 'identifier')
_PROTOCOL_PROPERTIES = frozenset(('XPT_RunId', 'XPT_Masked', 'XPT_Protocol'))


class PrivacyToolError(RuntimeError):
    """An operation failed without including workbook contents in its message."""


@dataclass(frozen=True)
class UnsupportedWorkbookFeatures:
    reasons: tuple[str, ...] = ()

    @property
    def has_unsupported(self):
        return bool(self.reasons)

    @property
    def pivot_table(self):
        return any('PivotTable' in label for label in self.reasons)

    def error_message(self, action_name='处理'):
        return f'检测到暂不支持的高级对象：{ "、".join(self.reasons) }。本次{action_name}已停止，原文件未修改。'


def resolve_path(path):
    return (PROJECT_ROOT / Path(path)).expanduser().resolve()


def is_within(child, parent):
    return Path(child).resolve().is_relative_to(Path(parent).resolve())


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')


def require_fernet():
    return Fernet


def check_not_overwrite_source(output_path, source_path):
    output, source = Path(output_path).resolve(), Path(source_path).resolve()
    if output == source or (output.exists() and source.exists() and os.path.samefile(output, source)):
        raise PrivacyToolError('Refusing to overwrite source workbook')


def get_zip_part_names(source):
    stream = io.BytesIO(source) if isinstance(source, bytes) else source
    position = stream.tell() if hasattr(stream, 'tell') else None
    try:
        if position is not None:
            stream.seek(0)
        with zipfile.ZipFile(stream) as package:
            return package.namelist()
    except Exception:
        raise PrivacyToolError('无法验证工作簿文件包；仅支持未加密的 .xlsx 文件。') from None
    finally:
        if position is not None:
            stream.seek(position)


def detect_unsupported_features(source):
    if isinstance(source, (str, Path)) and Path(source).suffix.lower() != '.xlsx':
        raise PrivacyToolError('V1.0 仅支持 .xlsx；不支持 .xlsm、旧版或加密工作簿。')
    parts = [p.replace('\\', '/').lower() for p in get_zip_part_names(source)]
    checks = {
        '数据透视表 (PivotTable)': ('xl/pivottables/', 'xl/pivotcache/'),
        '切片器 (Slicer)': ('xl/slicers/', 'xl/slicercaches/'),
        '外部数据连接 / Power Query': ('xl/connections.xml', 'xl/externallinks/'),
        '自定义 XML': ('customxml/',),
        'VBA 宏': ('xl/vbaproject',),
        '嵌入对象 / ActiveX': ('xl/embeddings/', 'xl/activex/'),
    }
    blocked = tuple(label for label, prefixes in checks.items()
                    if any(p.startswith(prefixes) for p in parts))
    if not blocked and ('[content_types].xml' not in parts or 'xl/workbook.xml' not in parts):
        raise PrivacyToolError('工作簿文件包结构无效；操作已停止。')
    return UnsupportedWorkbookFeatures(blocked)


def assert_supported_workbook(source, action_name='处理'):
    assessment = detect_unsupported_features(source)
    if assessment.has_unsupported:
        raise PrivacyToolError(assessment.error_message(action_name))


def load_workbook_private(path, password_provider=None):
    # Password-encrypted input was never an accepted V1 GUI path. No prompt,
    # decryption fallback, or legacy CLI is distributed by this V1 implementation.
    try:
        if isinstance(path, (str, Path)):
            if Path(path).suffix.lower() != '.xlsx':
                raise PrivacyToolError('V1.0 仅支持 .xlsx；不支持 .xlsm。')
            payload = Path(path).read_bytes()
        else:
            path.seek(0)
            payload = path.read()
        assert_supported_workbook(payload)
        return load_workbook(io.BytesIO(payload), data_only=False, keep_links=False)
    except PrivacyToolError:
        raise
    except Exception:
        raise PrivacyToolError('无法读取此 .xlsx 工作簿，操作已停止。') from None


def sanitize_workbook_metadata(wb):
    try:
        for name in _METADATA_FIELDS:
            setattr(wb.properties, name, None)
        wb.properties.creator = 'Excel Privacy Tool'
        wb.properties.lastModifiedBy = 'Excel Privacy Tool'
        retained = CustomPropertyList()
        for prop in wb.custom_doc_props:
            if prop.name in _PROTOCOL_PROPERTIES:
                retained.append(prop)
        wb.custom_doc_props = retained
        if any(p.name not in _PROTOCOL_PROPERTIES for p in wb.custom_doc_props):
            raise ValueError('incomplete cleanup')
    except Exception:
        raise PrivacyToolError('METADATA_SANITIZATION_FAILED：无法确认元数据已清理，已停止导出。') from None


def _verify_serialized_metadata(payload):
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as package:
            core = ET.fromstring(package.read('docProps/core.xml'))
            for node in core:
                name = node.tag.rsplit('}', 1)[-1]
                if name in _METADATA_FIELDS:
                    expected = 'Excel Privacy Tool' if name in ('creator', 'lastModifiedBy') else ''
                    if (node.text or '') != expected:
                        raise ValueError('core property remains')
            if 'docProps/custom.xml' in package.namelist():
                custom = ET.fromstring(package.read('docProps/custom.xml'))
                if any(node.attrib.get('name') not in _PROTOCOL_PROPERTIES for node in custom):
                    raise ValueError('custom property remains')
    except Exception:
        raise PrivacyToolError('METADATA_EXPORT_AUDIT_FAILED：文件包元数据检查未通过，未生成输出文件。') from None


def save_workbook_checked(wb, destination):
    # Serialize and inspect before creating a visible output. Exclusive creation
    # prevents overwriting an original or a concurrently created file.
    sanitize_workbook_metadata(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    payload = buffer.getvalue()
    _verify_serialized_metadata(payload)
    target = Path(destination)
    with target.open('xb') as stream:
        stream.write(payload)
        stream.flush()
        os.fsync(stream.fileno())


def current_checksums():
    return {p.relative_to(PROJECT_ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in LOCKED_FILES}


def verify_self_before_sensitive():
    try:
        expected = json.loads(APPROVED_CHECKSUMS.read_text(encoding='utf-8'))
        actual = current_checksums()
        if not expected or expected != actual:
            raise ValueError('integrity mismatch')
    except Exception:
        raise PrivacyToolError('INTEGRITY_CHECK_FAILED：程序文件完整性检查未通过，请使用原始发行包。') from None


def mask_workbook_v1(
    input_path: Path | str,
    output_dir: Path | str | None = None,
    private_key_dir: Path | str | None = None,
    header_row: int = 1,
    text_columns: list[str] | None = None,
    numeric_columns: list[str] | None = None,
    record_key: list[str] | str | None = None,
    target_sheet: str | None = None,
    numeric_scale_factor: Any = None,
    sheet_scopes: list[Any] | None = None,
) -> dict[str, Any]:
    """Execute canonical Checkpoint 4/V1 production masking pipeline.

    - Resolves TableScope (supporting header_row != 1)
    - Allocates cryptographically secure run_id and field tokens via TokenAllocator
    - Scales numeric columns via Analysis-Mode Numeric Masking (positive ratio k)
    - Captures Record Identity snapshots for accurate, tamper-proof recovery
    - Preserves formulas untouched, number formatting, and unselected columns
    - Activates OOXML recalculation properties (fullCalcOnLoad / forceFullCalc)
    - Sanitizes author metadata and stamps XPT properties into custom doc props
    - Encrypts complete RecoveryLedger with Fernet and writes atomically to private_keys
    """
    from decimal import Decimal
    from privacy_handler.identity import (
        TokenAllocator,
        get_ledger_paths,
        stamp_masked_workbook_identity,
    )
    from privacy_handler.numeric_masking import (
        NumericProtectedColumnConfig,
        ensure_workbook_recalculation_properties,
        generate_scale_factor,
        mask_numeric_table_scope,
    )
    from privacy_handler.recovery_ledger import (
        CURRENT_PROTOCOL_VERSION,
        LedgerConfig,
        LedgerEntry,
        LedgerRunInfo,
        ProtectedField,
        RecoveryLedger,
        encode_typed_value,
        save_encrypted_ledger,
    )
    from privacy_handler.key_detector import is_empty_value, is_formula_value
    from privacy_handler.table_scope import (
        extract_table_scope,
        resolve_safe_output_path,
        validate_composite_key,
    )

    resolved_input = Path(input_path).resolve()
    assert_supported_workbook(resolved_input, action_name="脱敏")

    resolved_out_dir = Path(output_dir).resolve() if output_dir else resolve_path("safe_for_codex")
    resolved_priv_dir = Path(private_key_dir).resolve() if private_key_dir else resolve_path("private_keys")

    if (is_within(resolved_priv_dir, resolve_path("safe_for_codex"))
            or is_within(resolved_priv_dir, resolved_out_dir)
            or is_within(resolved_out_dir, resolved_priv_dir)):
        raise PrivacyToolError("Private key directory must not be inside safe_for_codex")

    wb = load_workbook_private(resolved_input)
    target_sheets = [target_sheet] if target_sheet else wb.sheetnames

    # Parse global record_key list as default fallback
    global_key_fields_list: list[str] | None = None
    if record_key:
        if isinstance(record_key, str):
            global_key_fields_list = [k.strip() for k in record_key.split(",") if k.strip()]
        else:
            global_key_fields_list = [str(k).strip() for k in record_key if str(k).strip()]

    allocator = TokenAllocator()
    run_id = allocator.run_id

    fernet_cls = require_fernet()
    secret_bytes = fernet_cls.generate_key()
    fernet = fernet_cls(secret_bytes)

    ledger_enc_path, secret_key_path = get_ledger_paths(resolved_priv_dir, run_id)
    secret_key_path.parent.mkdir(parents=True, exist_ok=True)
    secret_key_path.write_bytes(secret_bytes)

    all_protected_fields: list[ProtectedField] = []
    all_entries: list[LedgerEntry] = []
    sheet_scopes_meta: list[dict[str, Any]] = []
    pending_record_snapshots: list[tuple[str, list[tuple[LedgerEntry, int]], list[tuple[str, int]]]] = []

    # Global text / numeric columns sets as defaults
    global_text_cols = {c.strip() for c in (text_columns or []) if c.strip()}
    global_num_cols = {c.strip() for c in (numeric_columns or []) if c.strip()}

    # Scale factor mapping by numeric analysis group
    group_scale_factors: dict[str, Decimal] = {}
    default_scale_factor = Decimal(str(numeric_scale_factor)) if numeric_scale_factor else None

    # Normalize sheet_scopes list if provided
    sheet_scope_map: dict[str, Any] = {}
    if sheet_scopes:
        target_sheets = []
        for sc in sheet_scopes:
            s_name = sc.sheet_name if hasattr(sc, "sheet_name") else sc.get("sheet_name")
            if s_name:
                sheet_scope_map[s_name] = sc
                target_sheets.append(s_name)
    else:
        target_sheets = [target_sheet] if target_sheet else wb.sheetnames

    for sheet_name in target_sheets:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]

        # Determine per-sheet configuration
        sc_obj = sheet_scope_map.get(sheet_name)
        if sc_obj is not None:
            cur_header_row = sc_obj.header_row if hasattr(sc_obj, "header_row") else sc_obj.get("header_row", header_row)
            cur_data_start = sc_obj.data_start_row if hasattr(sc_obj, "data_start_row") else sc_obj.get("data_start_row")
            raw_tcols = sc_obj.text_columns if hasattr(sc_obj, "text_columns") else sc_obj.get("text_columns", [])
            raw_ncols = sc_obj.numeric_columns if hasattr(sc_obj, "numeric_columns") else sc_obj.get("numeric_columns", [])
            raw_rkey = sc_obj.record_key if hasattr(sc_obj, "record_key") else sc_obj.get("record_key")
            group_id = sc_obj.numeric_analysis_group if hasattr(sc_obj, "numeric_analysis_group") else sc_obj.get("numeric_analysis_group")

            sheet_text_cols = [c.strip() for c in (raw_tcols or []) if c.strip()]
            sheet_num_cols = [c.strip() for c in (raw_ncols or []) if c.strip()]

            sheet_key_list: list[str] | None = None
            if raw_rkey:
                if isinstance(raw_rkey, str):
                    sheet_key_list = [k.strip() for k in raw_rkey.split(",") if k.strip()]
                else:
                    sheet_key_list = [str(k).strip() for k in raw_rkey if str(k).strip()]
            else:
                sheet_key_list = global_key_fields_list

            col_groups_map = sc_obj.numeric_column_groups if hasattr(sc_obj, "numeric_column_groups") else sc_obj.get("numeric_column_groups", {})
            if not isinstance(col_groups_map, dict):
                col_groups_map = {}
        else:
            cur_header_row = header_row
            cur_data_start = None
            sheet_text_cols = global_text_cols
            sheet_num_cols = global_num_cols
            sheet_key_list = global_key_fields_list
            group_id = None
            col_groups_map = {}

        scope = extract_table_scope(ws, header_row=cur_header_row, data_start_row=cur_data_start)
        sheet_scopes_meta.append(scope.to_dict())

        if scope.data_end_row is None or scope.data_end_row < scope.data_start_row:
            continue

        # Column letter/header lookup
        col_by_header = {col.header: col for col in scope.selected_columns}
        col_by_letter = {col.column_letter: col for col in scope.selected_columns}
        sheet_entry_rows: list[tuple[LedgerEntry, int]] = []

        # 1. Process Text Tokenizer columns
        for col_name in sheet_text_cols:
            col_desc = col_by_letter.get(col_name) or col_by_header.get(col_name)
            if not col_desc:
                continue

            c_idx = col_desc.column_index
            prot_field_id = f"{sheet_name}:{col_desc.column_letter}"
            all_protected_fields.append(
                ProtectedField(
                    logical_field=col_desc.header,
                    protection_type="text_token",
                    protected_field_id=prot_field_id,
                    sheet_name=sheet_name,
                    source_header_row=scope.header_row,
                    source_column_letter=col_desc.column_letter,
                    raw_header=col_desc.raw_header,
                    display_header=col_desc.display_header,
                )
            )

            for r in range(scope.data_start_row, scope.data_end_row + 1):
                cell = ws.cell(row=r, column=c_idx)
                orig_val = cell.value
                if orig_val is None or is_empty_value(orig_val) or is_formula_value(orig_val):
                    continue

                masked_token = allocator.allocate(col_desc.header, orig_val)
                cell.value = masked_token

                orig_tag, orig_ser = encode_typed_value(orig_val)
                entry = LedgerEntry(
                    sheet=sheet_name,
                    logical_field=col_desc.header,
                    key_masked_value=None,
                    original_type=orig_tag,
                    original_value=orig_ser,
                    masked_type="str",
                    masked_value=masked_token,
                    protection_type="text_token",
                    protected_field_id=prot_field_id,
                )
                all_entries.append(entry)
                sheet_entry_rows.append((entry, r))

        # 2. Process Numeric columns grouped by effective_group_id
        # Build NumericProtectedColumnConfig and map each to effective_group_id
        # Key priority:
        # 1. col_desc.column_letter in col_groups_map
        # 2. legacy sheet-level group_id
        # 3. None (uses default_scale_factor)
        configs_by_group: dict[str | None, list[NumericProtectedColumnConfig]] = {}
        for col_name in sheet_num_cols:
            col_desc = col_by_letter.get(col_name) or col_by_header.get(col_name)
            if not col_desc:
                continue
            prot_field_id = f"{sheet_name}:{col_desc.column_letter}"
            eff_gid = col_groups_map.get(col_desc.column_letter) or group_id

            num_cfg = NumericProtectedColumnConfig(
                sheet_name=sheet_name,
                column_letter=col_desc.column_letter,
                logical_field=col_desc.header,
                protected_field_id=prot_field_id,
                display_header=col_desc.display_header,
                raw_header=col_desc.raw_header,
            )
            configs_by_group.setdefault(eff_gid, []).append(num_cfg)

            all_protected_fields.append(
                ProtectedField(
                    logical_field=col_desc.header,
                    protection_type="numeric_scaled",
                    protected_field_id=prot_field_id,
                    sheet_name=sheet_name,
                    source_header_row=scope.header_row,
                    source_column_letter=col_desc.column_letter,
                    raw_header=col_desc.raw_header,
                    display_header=col_desc.display_header,
                    numeric_analysis_group=eff_gid,
                )
            )

        # Dispatch numeric masking per effective group
        for gid, grp_configs in configs_by_group.items():
            if gid is not None:
                if gid not in group_scale_factors:
                    group_scale_factors[gid] = default_scale_factor or generate_scale_factor()
                effective_k = group_scale_factors[gid]
            else:
                if default_scale_factor is None:
                    default_scale_factor = generate_scale_factor()
                effective_k = default_scale_factor

            numeric_entry_rows: list[int] = []
            num_entries, _ = mask_numeric_table_scope(
                ws=ws,
                scope=scope,
                numeric_configs=grp_configs,
                scale_factor=effective_k,
                key_fields=sheet_key_list,
                token_allocator=allocator,
                entry_row_indices=numeric_entry_rows,
            )
            if len(num_entries) != len(numeric_entry_rows):
                raise PrivacyToolError("Numeric ledger entry row tracking failed")
            all_entries.extend(num_entries)
            sheet_entry_rows.extend(zip(num_entries, numeric_entry_rows))

        # Keep each entry's source row only until the saved masked workbook can
        # supply the exact key representation that Compare will later read.
        if sheet_key_list:
            key_columns = [
                (key_name, (col_by_letter.get(key_name) or col_by_header.get(key_name)).column_index)
                for key_name in sheet_key_list
                if col_by_letter.get(key_name) or col_by_header.get(key_name)
            ]
            pending_record_snapshots.append((sheet_name, sheet_entry_rows, key_columns))

    # Calculation properties and metadata hygiene
    ensure_workbook_recalculation_properties(wb)
    stamp_masked_workbook_identity(wb, run_id=run_id, protocol_version=CURRENT_PROTOCOL_VERSION)
    sanitize_workbook_metadata(wb)

    # Resolve safe output workbook path
    output_suffix = ".masked.xlsx"
    masked_wb_path = resolve_safe_output_path(resolved_out_dir, resolved_input, suffix=output_suffix)
    save_workbook_checked(wb, masked_wb_path)

    # Excel serialization can normalize numeric key types (for example 153.0
    # becomes 153). Capture the Ledger identity from the persisted workbook,
    # not from the mutable in-memory sheet or the original plaintext key.
    saved_masked_wb = load_workbook_private(masked_wb_path)
    try:
        for sheet_name, entry_rows, key_columns in pending_record_snapshots:
            saved_ws = saved_masked_wb[sheet_name]
            snapshots_by_row: dict[int, list[dict[str, Any]]] = {}
            for entry, row_idx in entry_rows:
                if row_idx not in snapshots_by_row:
                    snapshot = []
                    for key_name, key_col_idx in key_columns:
                        key_value = saved_ws.cell(row=row_idx, column=key_col_idx).value
                        snapshot.append({
                            "field": key_name,
                            "type": type(key_value).__name__,
                            "value": key_value,
                        })
                    snapshots_by_row[row_idx] = snapshot
                entry.record_identity_masked = [dict(item) for item in snapshots_by_row[row_idx]]
    finally:
        saved_masked_wb.close()

    # Build and persist recovery ledger
    run_info = LedgerRunInfo(
        run_id=run_id,
        created_at=now_iso(),
        workbook_identity=run_id,
        source_extension=resolved_input.suffix or ".xlsx",
    )
    cfg_obj = LedgerConfig(
        confirmed_key_field="+".join(global_key_fields_list) if global_key_fields_list else None,
        numeric_mode="Mode_A_Analysis_Ratio" if any(pf.protection_type == "numeric_scaled" for pf in all_protected_fields) else None,
        sheet_scopes=sheet_scopes_meta,
    )
    ledger = RecoveryLedger(
        protocol_version=CURRENT_PROTOCOL_VERSION,
        run=run_info,
        configuration=cfg_obj,
        protected_fields=all_protected_fields,
        entries=all_entries,
    )
    save_encrypted_ledger(ledger, ledger_enc_path, secret_bytes)

    # Generate user-facing AI processing prompt
    from privacy_handler.prompt_generator import (
        build_prompt_context_from_metadata,
        generate_ai_processing_prompt,
    )
    prompt_ctx = build_prompt_context_from_metadata(
        sheet_scopes_meta=sheet_scopes_meta,
        protected_fields=all_protected_fields,
        confirmed_key_fields=cfg_obj.confirmed_key_field,
    )
    ai_prompt = generate_ai_processing_prompt(
        prompt_ctx,
        forbidden_values_to_check=[run_id, str(secret_bytes)],
    )

    return {
        "run_id": run_id,
        "masked_workbook_path": masked_wb_path,
        "ledger_enc_path": ledger_enc_path,
        "secret_key_path": secret_key_path,
        "ledger": ledger,
        "ai_prompt": ai_prompt,
    }

def main():
    if sys.argv[1:] != ['verify-self']:
        print('V1.0 uses privacy_app.py; supported diagnostic: verify-self', file=sys.stderr)
        return 2
    try:
        verify_self_before_sensitive()
    except PrivacyToolError as problem:
        print(str(problem), file=sys.stderr)
        return 1
    print('Privacy handler verification passed')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
