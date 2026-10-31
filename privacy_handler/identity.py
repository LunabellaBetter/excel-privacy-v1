"""Run identity, FP8 fingerprint, field-scoped TokenAllocator, and workbook XPT properties.

Provides unambiguous task identification, deterministic FP8 auxiliary fingerprinting,
CSPRNG token generation, field-scoped cross-sheet token allocation, and openpyxl
workbook custom properties stamping and reading.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
import re
import secrets
from typing import Any
import uuid

try:
    import openpyxl
    from openpyxl.packaging.custom import CustomPropertyList, StringProperty
except Exception:  # pragma: no cover
    openpyxl = None
    CustomPropertyList = None
    StringProperty = None


TOKEN_PREFIX = "MSK"
TOKEN_CHARSET = "ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789"
TOKEN_LENGTH = 8
FP8_LENGTH = 8

XPT_PROP_RUN_ID = "XPT_RunId"
XPT_PROP_MASKED = "XPT_Masked"
XPT_PROP_PROTOCOL = "XPT_Protocol"
XPT_ALLOWED_PROPS = {XPT_PROP_RUN_ID, XPT_PROP_MASKED, XPT_PROP_PROTOCOL}
DEFAULT_PROTOCOL_VERSION = "1.0"

TOKEN_REGEX = re.compile(r"^MSK-([A-F0-9]{8})-([A-Z0-9]{8})$")
TOKEN_SEARCH_REGEX = re.compile(r"MSK-([A-F0-9]{8})-([A-Z0-9]{8})")


class IdentityError(Exception):
    """Base error for identity operations."""


class AmbiguousRunIdentityError(IdentityError):
    """Raised when an FP8 fingerprint matches multiple recovery ledgers."""


def generate_run_id() -> str:
    """Generate standard UUIDv4 string for task identity."""
    return str(uuid.uuid4())


def compute_run_fingerprint(run_id: str) -> str:
    """Compute 8-character uppercase hex fingerprint (FP8) from run_id: SHA256(run_id)[:8].upper()."""
    if not run_id or not isinstance(run_id, str):
        raise IdentityError("run_id must be a non-empty string")
    return hashlib.sha256(run_id.encode("utf-8")).hexdigest()[:8].upper()


def generate_token(run_fingerprint: str, used_tokens: set[str] | None = None) -> str:
    """Generate random token MSK-[RUN_FP8]-[TOKEN8] using CSPRNG with collision avoidance."""
    if not run_fingerprint or len(run_fingerprint) != FP8_LENGTH:
        raise IdentityError(f"run_fingerprint must be {FP8_LENGTH} characters")

    while True:
        token_suffix = "".join(secrets.choice(TOKEN_CHARSET) for _ in range(TOKEN_LENGTH))
        token = f"{TOKEN_PREFIX}-{run_fingerprint}-{token_suffix}"
        if used_tokens is None or token not in used_tokens:
            if used_tokens is not None:
                used_tokens.add(token)
            return token


def parse_token(token: str) -> tuple[str, str] | None:
    """Parse a token into (run_fingerprint, token_suffix) if valid, else None."""
    m = TOKEN_REGEX.match(token)
    if m:
        return m.group(1), m.group(2)
    return None


class TokenAllocator:
    """Field-scoped token allocator managing (logical_field, original_value) mappings within a run."""

    def __init__(self, run_id: str | None = None) -> None:
        self.run_id: str = run_id if run_id else generate_run_id()
        self.run_fingerprint: str = compute_run_fingerprint(self.run_id)
        self._used_tokens: set[str] = set()
        self._field_links: dict[str, str] = {}
        self._mappings: dict[tuple[str, Any], str] = {}

    def link_field(self, alias_field: str, canonical_field: str) -> None:
        """Explicitly link an alias/derived field (e.g. 'Key', 'Key2') to a canonical field (e.g. '员工工号')."""
        self._field_links[alias_field] = canonical_field

    def resolve_field(self, field_name: str) -> str:
        """Resolve alias field name to its canonical field name."""
        return self._field_links.get(field_name, field_name)

    def allocate(self, logical_field: str, original_value: Any) -> str:
        """Allocate or retrieve consistent token for (logical_field, original_value)."""
        canonical_field = self.resolve_field(logical_field)
        key = (canonical_field, original_value)
        if key in self._mappings:
            return self._mappings[key]

        token = generate_token(self.run_fingerprint, self._used_tokens)
        self._mappings[key] = token
        return token

    def get_mappings(self) -> dict[tuple[str, Any], str]:
        """Return a copy of current field-value to token mappings."""
        return dict(self._mappings)


def stamp_masked_workbook_identity(
    wb: Any,
    run_id: str,
    protocol_version: str = DEFAULT_PROTOCOL_VERSION,
) -> None:
    """Stamp non-secret XPT properties into openpyxl Workbook custom doc props."""
    from openpyxl.packaging.custom import CustomPropertyList, StringProperty

    if not hasattr(wb, "custom_doc_props") or wb.custom_doc_props is None:
        wb.custom_doc_props = CustomPropertyList()

    new_props = CustomPropertyList()
    for prop in wb.custom_doc_props:
        if getattr(prop, "name", None) not in XPT_ALLOWED_PROPS:
            new_props.append(prop)

    new_props.append(StringProperty(name=XPT_PROP_RUN_ID, value=str(run_id)))
    new_props.append(StringProperty(name=XPT_PROP_MASKED, value="true"))
    new_props.append(StringProperty(name=XPT_PROP_PROTOCOL, value=str(protocol_version)))
    wb.custom_doc_props = new_props


def read_workbook_identity(wb_or_path: Any) -> dict[str, str | None]:
    """Read XPT identity properties from an openpyxl Workbook or file path."""
    result: dict[str, str | None] = {
        "run_id": None,
        "is_masked": None,
        "protocol_version": None,
    }
    close_needed = False
    if isinstance(wb_or_path, (str, Path)):
        import openpyxl
        wb = openpyxl.load_workbook(wb_or_path, read_only=True)
        close_needed = True
    else:
        wb = wb_or_path

    try:
        if hasattr(wb, "custom_doc_props") and wb.custom_doc_props:
            for prop in wb.custom_doc_props:
                name = getattr(prop, "name", None)
                val = getattr(prop, "value", None)
                if name == XPT_PROP_RUN_ID:
                    result["run_id"] = str(val) if val is not None else None
                elif name == XPT_PROP_MASKED:
                    result["is_masked"] = str(val) if val is not None else None
                elif name == XPT_PROP_PROTOCOL:
                    result["protocol_version"] = str(val) if val is not None else None
    finally:
        if close_needed and hasattr(wb, "close"):
            wb.close()
    return result


def is_masked_workbook(wb_or_path: Any) -> bool:
    """Determine if a workbook has been masked by checking XPT props or token existence."""
    # Check A: XPT custom properties
    ident = read_workbook_identity(wb_or_path)
    if ident["is_masked"] == "true" or ident["run_id"] is not None:
        return True

    # Check B: Scan cells for MSK-[FP8]-[TOKEN8] pattern
    close_needed = False
    if isinstance(wb_or_path, (str, Path)):
        import openpyxl
        wb = openpyxl.load_workbook(wb_or_path, read_only=True, data_only=True)
        close_needed = True
    else:
        wb = wb_or_path

    try:
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            for row in ws.iter_rows(values_only=True):
                for val in row:
                    if val and isinstance(val, str) and TOKEN_SEARCH_REGEX.search(val):
                        return True
    finally:
        if close_needed and hasattr(wb, "close"):
            wb.close()
    return False


def get_ledger_paths(private_key_dir: Path | str, run_id: str) -> tuple[Path, Path]:
    """Return (ledger_enc_path, secret_key_path) bound to the given run_id."""
    dir_path = Path(private_key_dir)
    return (
        dir_path / f"{run_id}.recovery_ledger.enc",
        dir_path / f"{run_id}.mask_secret.key",
    )


def resolve_ledger_by_fingerprint(
    private_key_dir: Path | str,
    fp8: str,
    candidate_files: list[Path] | None = None,
) -> Path:
    """Resolve recovery ledger by FP8 fallback clue. Refuses to guess on ambiguity."""
    fp8_target = fp8.strip().upper()
    if len(fp8_target) != FP8_LENGTH:
        raise IdentityError(f"Invalid FP8 length: expected {FP8_LENGTH}, got {len(fp8_target)}")

    if candidate_files is not None:
        candidates = candidate_files
    else:
        dir_path = Path(private_key_dir)
        if not dir_path.exists():
            raise FileNotFoundError(f"Directory not found: {dir_path}")
        candidates = list(dir_path.glob("*.recovery_ledger.enc"))

    matched: list[Path] = []
    for c in candidates:
        filename = c.name
        if filename.endswith(".recovery_ledger.enc"):
            candidate_run_id = filename[:-len(".recovery_ledger.enc")]
            try:
                candidate_fp8 = compute_run_fingerprint(candidate_run_id)
                if candidate_fp8 == fp8_target:
                    matched.append(c)
            except Exception:
                continue

    if len(matched) == 0:
        raise FileNotFoundError(f"No recovery ledger found matching FP8 fingerprint '{fp8_target}'")
    if len(matched) > 1:
        names = [m.name for m in matched]
        raise AmbiguousRunIdentityError(
            f"Ambiguous run fingerprint '{fp8_target}': matches multiple recovery ledgers: {names}"
        )
    return matched[0]
